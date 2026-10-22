"""
test_excel_export.py
---------------------
Dedicated automated test suite for the Excel Attendance Download feature:
1. Monthly report generation (month + year).
2. Custom date-range report generation.
3. Full-semester report generation (Odd/Even semester).
4. Empty attendance data (valid .xlsx generated, 0% attendance, no false absences).
5. Invalid date ranges and formats (start_date > end_date, malformed strings).
6. Student details, session dates, P/A markings, and calculation accuracy.
7. Genuine .xlsx binary format and sheet content verification via openpyxl.
8. RBAC security enforcement: Teacher (200), Admin (200), Student (403), Security (403), Anon (401).
9. Confirmation that security-gate observations are strictly excluded from classroom Excel export.
"""

from datetime import datetime
import io
from pathlib import Path
import sys

backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

from fastapi.testclient import TestClient
import openpyxl
import pytest

from app.auth import create_access_token
from app.db.database import get_db, init_db
from app.main import app
from app.recognition_tracker import reset_tracker

client = TestClient(app)


@pytest.fixture(autouse=True)
def clean_test_environment():
    """Ensure clean tracker and database state for each test."""
    init_db()
    reset_tracker()
    yield
    reset_tracker()


def get_token_headers(username: str, role: str) -> dict:
    """Generate authenticated bearer header for test user."""
    with get_db() as conn:
        user = conn.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()
        assert user, f"User '{username}' must exist in database"
        user_id = user["id"]
    token = create_access_token({"sub": str(user_id), "role": role})
    return {"Authorization": f"Bearer {token}"}


# ── 1. RBAC & Access Control ──────────────────────────────────────────────────

def test_excel_export_unauthenticated_fails():
    """Unauthenticated requests must be rejected with HTTP 401."""
    resp = client.get("/api/attendance/export/excel?report_type=monthly")
    assert resp.status_code == 401


def test_excel_export_student_forbidden():
    """Student role must be forbidden with HTTP 403."""
    stu_headers = get_token_headers("prapti", "student")
    resp = client.get("/api/attendance/export/excel?report_type=monthly", headers=stu_headers)
    assert resp.status_code == 403
    assert "Forbidden" in resp.text or "Not enough permissions" in resp.text or resp.status_code == 403


def test_excel_export_security_forbidden():
    """Security role must be forbidden with HTTP 403."""
    sec_headers = get_token_headers("security", "security")
    resp = client.get("/api/attendance/export/excel?report_type=monthly", headers=sec_headers)
    assert resp.status_code == 403


def test_excel_export_teacher_and_admin_allowed():
    """Teacher and admin roles must be authorized (HTTP 200)."""
    tch_headers = get_token_headers("teacher", "teacher")
    resp_tch = client.get("/api/attendance/export/excel?report_type=monthly", headers=tch_headers)
    assert resp_tch.status_code == 200
    assert resp_tch.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    adm_headers = get_token_headers("admin", "admin")
    resp_adm = client.get("/api/attendance/export/excel?report_type=monthly", headers=adm_headers)
    assert resp_adm.status_code == 200
    assert resp_adm.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── 2. Input Validation & Error Handling ─────────────────────────────────────

def test_invalid_report_type_rejected():
    """Invalid report_type returns 400 Bad Request."""
    tch_headers = get_token_headers("teacher", "teacher")
    resp = client.get("/api/attendance/export/excel?report_type=annual", headers=tch_headers)
    assert resp.status_code == 400
    assert "Unsupported report_type" in resp.json()["detail"]


def test_invalid_month_format_rejected():
    """Malformed month string returns 400 Bad Request."""
    tch_headers = get_token_headers("teacher", "teacher")
    for bad_month in ["2026-13", "invalid-month", "2026/10", "2026"]:
        resp = client.get(f"/api/attendance/export/excel?report_type=monthly&month={bad_month}", headers=tch_headers)
        assert resp.status_code == 400
        assert "Invalid month format" in resp.json()["detail"]


def test_custom_date_range_invalid_order_rejected():
    """start_date > end_date returns 400 Bad Request."""
    tch_headers = get_token_headers("teacher", "teacher")
    resp = client.get(
        "/api/attendance/export/excel?report_type=custom&start_date=2026-10-25&end_date=2026-10-10",
        headers=tch_headers,
    )
    assert resp.status_code == 400
    assert "Start date cannot be after end date" in resp.json()["detail"]


def test_custom_date_range_malformed_date_rejected():
    """Malformed start_date or end_date returns 400 Bad Request."""
    tch_headers = get_token_headers("teacher", "teacher")
    resp = client.get(
        "/api/attendance/export/excel?report_type=custom&start_date=not-a-date&end_date=2026-10-10",
        headers=tch_headers,
    )
    assert resp.status_code == 400
    assert "Invalid date format" in resp.json()["detail"]


# ── 3. Monthly Report & Calculation Accuracy ─────────────────────────────────

def test_monthly_report_calculations_and_content():
    """
    Verify monthly report:
    - Lists registered students
    - Has correct session date columns
    - Accurately marks Present (P) and Absent (A)
    - Accurately calculates Total Classes Held, Present, Absent, and Attendance %
    """
    tch_headers = get_token_headers("teacher", "teacher")

    with get_db() as conn:
        students = conn.execute("SELECT roll_no, name FROM students ORDER BY roll_no ASC").fetchall()
        assert len(students) >= 2, "Test requires at least 2 students"
        s1_roll = students[0]["roll_no"]
        s2_roll = students[1]["roll_no"]

        # Clean existing records for test month
        conn.execute("DELETE FROM attendance WHERE date LIKE '2026-10%'")

        # Class date 1: Both students present
        conn.execute(
            "INSERT INTO attendance (roll_no, date, entry_time, status, similarity) VALUES (?, '2026-10-05', '09:00:00', 'present', 0.90)",
            (s1_roll,),
        )
        conn.execute(
            "INSERT INTO attendance (roll_no, date, entry_time, status, similarity) VALUES (?, '2026-10-05', '09:02:00', 'present', 0.92)",
            (s2_roll,),
        )

        # Class date 2: Only student 1 present (student 2 absent)
        conn.execute(
            "INSERT INTO attendance (roll_no, date, entry_time, status, similarity) VALUES (?, '2026-10-06', '09:05:00', 'present', 0.88)",
            (s1_roll,),
        )

    resp = client.get(
        "/api/attendance/export/excel?report_type=monthly&month=2026-10",
        headers=tch_headers,
    )
    assert resp.status_code == 200
    assert "attachment" in resp.headers["content-disposition"]
    assert "SmartAttendance_monthly_2026-10-01_to_2026-10-31.xlsx" in resp.headers["content-disposition"]

    # Verify openpyxl can parse the file
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    assert ws.title == "Classroom Attendance"

    # Banner checks
    assert "SMARTATTENDANCE" in str(ws.cell(row=1, column=1).value)
    assert "Month: October 2026" in str(ws.cell(row=2, column=1).value)
    assert "Total Conducted Sessions: 2" in str(ws.cell(row=2, column=1).value)
    assert "Generated On:" in str(ws.cell(row=3, column=1).value)
    assert "Gate Observations Excluded" in str(ws.cell(row=3, column=1).value)

    # Column header checks
    headers = [ws.cell(row=5, column=col).value for col in range(1, ws.max_column + 1)]
    assert "Roll No" in headers
    assert "Student Name" in headers
    assert "10/05" in headers
    assert "10/06" in headers
    assert "Present" in headers
    assert "Absent" in headers
    assert "Total Classes Held" in headers
    assert "Attendance %" in headers

    # Row data checks
    s1_row: int = 0
    s2_row: int = 0
    for r in range(6, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == s1_roll:
            s1_row = r
        elif ws.cell(row=r, column=1).value == s2_roll:
            s2_row = r

    assert s1_row > 0
    assert s2_row > 0

    col_10_05 = headers.index("10/05") + 1
    col_10_06 = headers.index("10/06") + 1
    col_present = headers.index("Present") + 1
    col_absent = headers.index("Absent") + 1
    col_held = headers.index("Total Classes Held") + 1
    col_pct = headers.index("Attendance %") + 1

    # Student 1: Present on both dates (2 held, 2 attended, 0 absent, 100.0%)
    assert ws.cell(row=s1_row, column=col_10_05).value == "P"
    assert ws.cell(row=s1_row, column=col_10_06).value == "P"
    assert ws.cell(row=s1_row, column=col_present).value == 2
    assert ws.cell(row=s1_row, column=col_absent).value == 0
    assert ws.cell(row=s1_row, column=col_held).value == 2
    assert ws.cell(row=s1_row, column=col_pct).value == "100.0%"

    # Student 2: Present on 10/05, Absent on 10/06 (2 held, 1 attended, 1 absent, 50.0%)
    assert ws.cell(row=s2_row, column=col_10_05).value == "P"
    assert ws.cell(row=s2_row, column=col_10_06).value == "A"
    assert ws.cell(row=s2_row, column=col_present).value == 1
    assert ws.cell(row=s2_row, column=col_absent).value == 1
    assert ws.cell(row=s2_row, column=col_held).value == 2
    assert ws.cell(row=s2_row, column=col_pct).value == "50.0%"


# ── 4. Custom Date Range Report ──────────────────────────────────────────────

def test_custom_date_range_report():
    """Verify custom date-range export."""
    tch_headers = get_token_headers("teacher", "teacher")

    with get_db() as conn:
        stu = conn.execute("SELECT roll_no FROM students LIMIT 1").fetchone()
        s_roll = stu["roll_no"]
        conn.execute("DELETE FROM attendance WHERE date BETWEEN '2026-09-01' AND '2026-09-30'")
        conn.execute(
            "INSERT INTO attendance (roll_no, date, entry_time, status, similarity) VALUES (?, '2026-09-15', '10:00:00', 'present', 0.93)",
            (s_roll,),
        )

    resp = client.get(
        "/api/attendance/export/excel?report_type=custom&start_date=2026-09-01&end_date=2026-09-20",
        headers=tch_headers,
    )
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    assert "Period: 2026-09-01 to 2026-09-20" in str(ws.cell(row=2, column=1).value)
    assert "Total Conducted Sessions: 1" in str(ws.cell(row=2, column=1).value)


# ── 5. Full-Semester Report ──────────────────────────────────────────────────

def test_semester_report():
    """Verify full-semester report defaults or custom semester dates."""
    tch_headers = get_token_headers("teacher", "teacher")
    resp = client.get("/api/attendance/export/excel?report_type=semester", headers=tch_headers)
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None
    assert "Semester" in str(ws.cell(row=2, column=1).value)


# ── 6. Empty Attendance Data (No Classes Conducted) ──────────────────────────

def test_empty_attendance_data():
    """
    When no classes were held in the specified period:
    - Total Conducted Sessions is 0
    - Students have 0 Present, 0 Absent, 0 Held, 0.0%
    - Students are NOT falsely marked absent for unheld classes
    """
    tch_headers = get_token_headers("teacher", "teacher")

    resp = client.get(
        "/api/attendance/export/excel?report_type=custom&start_date=2025-01-01&end_date=2025-01-10",
        headers=tch_headers,
    )
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None

    assert "Total Conducted Sessions: 0" in str(ws.cell(row=2, column=1).value)
    headers = [ws.cell(row=5, column=col).value for col in range(1, ws.max_column + 1)]
    assert "Present" in headers
    assert "Absent" in headers
    assert "Total Classes Held" in headers
    assert "Attendance %" in headers

    # Verify first student has 0 present, 0 absent, 0 conducted, 0.0%
    col_present = headers.index("Present") + 1
    col_absent = headers.index("Absent") + 1
    col_held = headers.index("Total Classes Held") + 1
    col_pct = headers.index("Attendance %") + 1

    assert ws.cell(row=6, column=col_present).value == 0
    assert ws.cell(row=6, column=col_absent).value == 0
    assert ws.cell(row=6, column=col_held).value == 0
    assert ws.cell(row=6, column=col_pct).value == "0.0%"


# ── 7. Gate Observation Exclusion ────────────────────────────────────────────

def test_gate_observations_strictly_excluded_from_excel():
    """
    Verify that campus gate observations (in security_events) are NEVER counted
    as classroom sessions or attendance in the Excel export.
    """
    tch_headers = get_token_headers("teacher", "teacher")
    gate_date = "2026-11-15"

    with get_db() as conn:
        conn.execute("DELETE FROM attendance WHERE date = ?", (gate_date,))
        conn.execute("DELETE FROM security_events WHERE timestamp LIKE '2026-11-15%'")

        # Insert a gate observation event for Prapti
        conn.execute(
            """
            INSERT INTO security_events (
                event_id, event_type, name, identity_category, source_camera,
                similarity, resolution_status, timestamp, last_seen_at
            ) VALUES (
                'evt-gate-test-001', 'GATE_OBSERVATION', 'Prapti', 'STUDENT', 'gate_webcam',
                0.91, 'resolved', '2026-11-15 08:30:00', '2026-11-15 08:30:00'
            )
            """
        )

    # Request Excel report covering that date
    resp = client.get(
        f"/api/attendance/export/excel?report_type=custom&start_date={gate_date}&end_date={gate_date}",
        headers=tch_headers,
    )
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws is not None

    # Since no classroom attendance was marked on 2026-11-15, total conducted sessions must be 0!
    assert "Total Conducted Sessions: 0" in str(ws.cell(row=2, column=1).value)
    headers = [ws.cell(row=5, column=col).value for col in range(1, ws.max_column + 1)]
    assert "11/15" not in headers
    assert "2026-11-15" not in headers
